from __future__ import annotations

import csv
import mimetypes
import re
from pathlib import Path
from urllib.parse import urlencode

from markupsafe import Markup, escape
from openpyxl import load_workbook

from app.config import Settings
from app.types import SearchHit

try:
    import xlrd
except ImportError:  # pragma: no cover - optional dependency for legacy Excel files
    xlrd = None


def resolve_document_path(settings: Settings, raw_path: str) -> Path:
    candidate = Path(raw_path).expanduser()
    if not candidate.is_absolute():
        candidate = settings.documents_dir / candidate

    resolved = candidate.resolve(strict=True)
    documents_root = settings.documents_dir.resolve()
    resolved.relative_to(documents_root)
    if not resolved.is_file():
        raise FileNotFoundError(str(resolved))
    return resolved


def build_source_preview_url(metadata: dict[str, object]) -> str:
    params: dict[str, object] = {"path": str(metadata.get("source_path", ""))}
    for key in ("page", "sheet_name", "row_number"):
        value = metadata.get(key)
        if value is not None:
            params[key] = value
    return f"/documents/preview?{urlencode(params)}"


def build_document_file_url(path: Path) -> str:
    return f"/documents/file?{urlencode({'path': str(path)})}"


def build_native_open_url(path: Path) -> str:
    resolved_uri = path.resolve().as_uri()
    suffix = path.suffix.lower()
    if suffix in {".xls", ".xlsx", ".csv"}:
        return f"ms-excel:ofe|u|{resolved_uri}"
    if suffix == ".docx":
        return f"ms-word:ofe|u|{resolved_uri}"
    return build_document_file_url(path)


def render_answer_html(answer: str, sources: list[SearchHit]) -> Markup:
    pattern = re.compile(r"\[SOURCE\s+(\d+)\]")
    parts: list[Markup | str] = []
    last = 0

    for match in pattern.finditer(answer):
        parts.append(escape(answer[last:match.start()]))
        source_index = int(match.group(1)) - 1
        label = match.group(0)
        if 0 <= source_index < len(sources):
            href = build_source_preview_url(sources[source_index].metadata)
            parts.append(
                Markup(
                    f'<a class="source-inline-link" href="{escape(href)}" '
                    f'target="_blank" rel="noopener">{escape(label)}</a>'
                )
            )
        else:
            parts.append(escape(label))
        last = match.end()

    parts.append(escape(answer[last:]))
    html = Markup("").join(Markup(str(part)) for part in parts)
    return Markup(str(html).replace("\n", "<br>"))


def guess_media_type(path: Path) -> str:
    media_type, _ = mimetypes.guess_type(path.name)
    return media_type or "application/octet-stream"


def build_preview_context(
    path: Path,
    page: int | None = None,
    sheet_name: str | None = None,
    row_number: int | None = None,
) -> dict[str, object]:
    suffix = path.suffix.lower()
    file_url = build_document_file_url(path)
    native_open_url = build_native_open_url(path)
    context: dict[str, object] = {
        "file_name": path.name,
        "file_path": str(path),
        "file_url": file_url,
        "native_open_url": native_open_url,
        "kind": "file",
        "page": page,
        "sheet_name": sheet_name,
        "row_number": row_number,
        "table_headers": [],
        "table_rows": [],
    }

    if suffix == ".pdf":
        context["kind"] = "pdf"
        context["embed_url"] = f"{file_url}#page={page}" if page else file_url
        return context

    if suffix in {".xlsx", ".xls", ".csv"}:
        headers, rows, actual_sheet = build_tabular_preview(path, sheet_name=sheet_name, row_number=row_number)
        context["kind"] = "table"
        context["table_headers"] = headers
        context["table_rows"] = rows
        context["sheet_name"] = actual_sheet
        return context

    return context


def build_tabular_preview(
    path: Path,
    sheet_name: str | None = None,
    row_number: int | None = None,
    window_size: int = 3,
) -> tuple[list[str], list[dict[str, object]], str | None]:
    target_row = max(1, int(row_number or 1))
    row_start = max(1, target_row - window_size)
    row_end = target_row + window_size

    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _build_xlsx_preview(path, sheet_name, target_row, row_start, row_end)
    if suffix == ".xls":
        return _build_xls_preview(path, sheet_name, target_row, row_start, row_end)
    return _build_csv_preview(path, target_row, row_start, row_end)


def _build_xlsx_preview(
    path: Path,
    sheet_name: str | None,
    target_row: int,
    row_start: int,
    row_end: int,
) -> tuple[list[str], list[dict[str, object]], str | None]:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.worksheets[0]
        rows: list[dict[str, object]] = []
        max_columns = 0
        for current_row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if current_row_number < row_start:
                continue
            if current_row_number > row_end:
                break
            cells = [_format_preview_cell(value) for value in row]
            max_columns = max(max_columns, len(cells))
            rows.append(
                {
                    "row_number": current_row_number,
                    "cells": cells,
                    "is_target": current_row_number == target_row,
                }
            )
        headers = _build_column_headers(max_columns)
        _pad_rows(rows, max_columns)
        return headers, rows, sheet.title
    finally:
        workbook.close()


def _build_xls_preview(
    path: Path,
    sheet_name: str | None,
    target_row: int,
    row_start: int,
    row_end: int,
) -> tuple[list[str], list[dict[str, object]], str | None]:
    if xlrd is None:
        return [], [], sheet_name

    workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
    try:
        sheet = workbook.sheet_by_name(sheet_name) if sheet_name and sheet_name in workbook.sheet_names() else workbook.sheet_by_index(0)
        rows: list[dict[str, object]] = []
        max_columns = 0
        for current_row_number in range(row_start, min(row_end, sheet.nrows) + 1):
            cells = [
                _format_xls_preview_cell(workbook, cell)
                for cell in sheet.row(current_row_number - 1)
            ]
            max_columns = max(max_columns, len(cells))
            rows.append(
                {
                    "row_number": current_row_number,
                    "cells": cells,
                    "is_target": current_row_number == target_row,
                }
            )
        headers = _build_column_headers(max_columns)
        _pad_rows(rows, max_columns)
        return headers, rows, sheet.name
    finally:
        workbook.release_resources()


def _build_csv_preview(
    path: Path,
    target_row: int,
    row_start: int,
    row_end: int,
) -> tuple[list[str], list[dict[str, object]], str | None]:
    rows: list[dict[str, object]] = []
    max_columns = 0
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for current_row_number, row in enumerate(reader, start=1):
            if current_row_number < row_start:
                continue
            if current_row_number > row_end:
                break
            cells = [str(value).strip() for value in row]
            max_columns = max(max_columns, len(cells))
            rows.append(
                {
                    "row_number": current_row_number,
                    "cells": cells,
                    "is_target": current_row_number == target_row,
                }
            )

    headers = _build_column_headers(max_columns)
    _pad_rows(rows, max_columns)
    return headers, rows, "CSV"


def _build_column_headers(total_columns: int) -> list[str]:
    headers = []
    for index in range(total_columns):
        value = index + 1
        label = ""
        while value > 0:
            value, remainder = divmod(value - 1, 26)
            label = chr(65 + remainder) + label
        headers.append(label)
    return headers


def _pad_rows(rows: list[dict[str, object]], total_columns: int) -> None:
    for row in rows:
        cells = row["cells"]
        if isinstance(cells, list) and len(cells) < total_columns:
            cells.extend([""] * (total_columns - len(cells)))


def _format_preview_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _format_xls_preview_cell(workbook: object, cell: object) -> str:
    if xlrd is None:
        return ""
    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return ""
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            date_value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
            return date_value.isoformat(sep=" ", timespec="seconds")
        except (OverflowError, ValueError):
            return str(cell.value).strip()
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "TRUE" if bool(cell.value) else "FALSE"
    if cell.ctype == xlrd.XL_CELL_NUMBER:
        number = float(cell.value)
        return str(int(number)) if number.is_integer() else str(number)
    return str(cell.value).strip()
