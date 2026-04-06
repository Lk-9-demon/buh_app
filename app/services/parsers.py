from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from app.config import Settings
from app.types import ParsedChunk


class DocumentParser:
    def __init__(self, settings: Settings) -> None:
        self.settings = settings

    def supports(self, path: Path) -> bool:
        return path.suffix.lower() in self.settings.allowed_extensions

    def parse_file(self, path: Path) -> list[ParsedChunk]:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._parse_pdf(path)
        if suffix == ".xlsx":
            return self._parse_xlsx(path)
        if suffix == ".docx":
            return self._parse_docx(path)
        if suffix == ".csv":
            return self._parse_csv(path)
        if suffix == ".txt":
            return self._parse_txt(path)
        raise ValueError(f"Unsupported file format: {suffix}")

    def _parse_pdf(self, path: Path) -> list[ParsedChunk]:
        reader = PdfReader(str(path))
        chunks: list[ParsedChunk] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = self._extract_pdf_text(page)
            if not text:
                continue
            for chunk_index, chunk_text in enumerate(self._split_text(text)):
                chunks.append(
                    ParsedChunk(
                        text=chunk_text,
                        metadata={
                            "page": page_number,
                            "chunk_index": chunk_index,
                        },
                    )
                )
        return chunks

    def _extract_pdf_text(self, page: object) -> str:
        raw_text = ""
        try:
            raw_text = page.extract_text(extraction_mode="layout") or ""
        except TypeError:
            raw_text = page.extract_text() or ""
        return self._clean_text(raw_text)

    def _parse_docx(self, path: Path) -> list[ParsedChunk]:
        document = Document(str(path))
        blocks: list[str] = []

        for paragraph in document.paragraphs:
            text = self._clean_text(paragraph.text)
            if text:
                blocks.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [self._clean_text(cell.text) for cell in row.cells]
                row_text = " | ".join(cell for cell in cells if cell)
                if row_text:
                    blocks.append(row_text)

        chunks: list[ParsedChunk] = []
        for block_index, block in enumerate(blocks):
            for chunk_index, chunk_text in enumerate(self._split_text(block)):
                chunks.append(
                    ParsedChunk(
                        text=chunk_text,
                        metadata={
                            "block_index": block_index,
                            "chunk_index": chunk_index,
                        },
                    )
                )
        return chunks

    def _parse_xlsx(self, path: Path) -> list[ParsedChunk]:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        chunks: list[ParsedChunk] = []

        for sheet in workbook.worksheets:
            headers: list[str] | None = None
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                values = [self._format_cell(value) for value in row]
                if not any(values):
                    continue

                if headers is None and self._looks_like_header(values):
                    headers = [value if value else f"Column {index + 1}" for index, value in enumerate(values)]
                    continue

                labels = headers or [f"Column {index + 1}" for index in range(len(values))]
                parts = []
                for index, value in enumerate(values):
                    if not value:
                        continue
                    label = labels[index] if index < len(labels) else f"Column {index + 1}"
                    parts.append(f"{label}: {value}")

                if not parts:
                    continue

                chunks.append(
                    ParsedChunk(
                        text=" | ".join(parts),
                        metadata={
                            "sheet_name": sheet.title,
                            "row_number": row_number,
                            "chunk_index": 0,
                        },
                    )
                )

        return chunks

    def _parse_csv(self, path: Path) -> list[ParsedChunk]:
        chunks: list[ParsedChunk] = []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            rows = list(reader)

        headers: list[str] | None = None
        for row_number, row in enumerate(rows, start=1):
            values = [self._clean_text(value) for value in row]
            if not any(values):
                continue

            if headers is None and self._looks_like_header(values):
                headers = [value if value else f"Column {index + 1}" for index, value in enumerate(values)]
                continue

            labels = headers or [f"Column {index + 1}" for index in range(len(values))]
            parts = []
            for index, value in enumerate(values):
                if not value:
                    continue
                label = labels[index] if index < len(labels) else f"Column {index + 1}"
                parts.append(f"{label}: {value}")
            if parts:
                chunks.append(
                    ParsedChunk(
                        text=" | ".join(parts),
                        metadata={
                            "sheet_name": "CSV",
                            "row_number": row_number,
                            "chunk_index": 0,
                        },
                    )
                )
        return chunks

    def _parse_txt(self, path: Path) -> list[ParsedChunk]:
        raw_text = path.read_text(encoding="utf-8", errors="ignore")
        cleaned = self._clean_text(raw_text)
        return [
            ParsedChunk(text=chunk_text, metadata={"chunk_index": chunk_index})
            for chunk_index, chunk_text in enumerate(self._split_text(cleaned))
        ]

    def _split_text(self, text: str) -> list[str]:
        cleaned = self._clean_text(text)
        if not cleaned:
            return []
        if len(cleaned) <= self.settings.max_chunk_chars:
            return [cleaned]

        max_chars = self.settings.max_chunk_chars
        overlap = self.settings.chunk_overlap_chars
        chunks: list[str] = []
        start = 0

        while start < len(cleaned):
            end = min(len(cleaned), start + max_chars)
            if end < len(cleaned):
                split_at = cleaned.rfind(" ", start, end)
                if split_at > start + (max_chars // 2):
                    end = split_at

            chunk = cleaned[start:end].strip()
            if chunk:
                chunks.append(chunk)

            if end >= len(cleaned):
                break

            start = max(end - overlap, start + 1)

        return chunks

    def _clean_text(self, value: str) -> str:
        lines = []
        for line in value.splitlines():
            compact = re.sub(r"\s+", " ", line).strip()
            if compact:
                lines.append(compact)
        return "\n".join(lines)

    def _looks_like_header(self, values: Iterable[str]) -> bool:
        non_empty = [value for value in values if value]
        if len(non_empty) < 2:
            return False
        numeric_count = sum(1 for value in non_empty if self._is_numeric(value))
        return numeric_count / len(non_empty) < 0.4

    def _is_numeric(self, value: str) -> bool:
        normalized = value.replace(" ", "").replace(",", ".")
        try:
            float(normalized)
        except ValueError:
            return False
        return True

    def _format_cell(self, value: object) -> str:
        if value is None:
            return ""
        return self._clean_text(str(value))
